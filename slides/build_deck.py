#!/usr/bin/env python3
"""Build the FrontX deck by inserting our 5 slides into the existing template.

Reads compact slide markdown from `slides/deck/`, picks rich template layouts
(Three Content Blocks, 5 steps w icons, Four/Six Content Blocks, Timeline (5
years)), populates each layout's placeholders, and splices the slides into the
template at position 22 (replacing the placeholder slide).

The full content prototypes in `slides/01-*.md … 05-*.md` are untouched and
remain as references; the deck builds from `slides/deck/*.md`.
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

SLIDES_DIR = Path(__file__).parent
DECK_DIR = SLIDES_DIR / "deck"
TEMPLATE = SLIDES_DIR / "Cyber Console Presentation 28.04.2026.pptx"
OUTPUT = SLIDES_DIR / "frontx-deck.pptx"

PLACEHOLDER_SLIDE_NUMBER = 22  # 1-indexed slide that gets replaced


# --------------------------------------------------------------------------
# Markdown parsing
# --------------------------------------------------------------------------


def parse_compact(path: Path) -> dict:
    """Parse a compact slide file into title / layout / items / speaker_notes."""
    text = path.read_text()
    sections: dict = {}
    current = None
    buf: list = []
    for line in text.split("\n"):
        m = re.match(r"^##\s+(.+?)\s*$", line)
        if m:
            if current:
                sections[current] = "\n".join(buf).strip()
            current = m.group(1).strip()
            buf = []
        elif current is not None:
            buf.append(line)
    if current:
        sections[current] = "\n".join(buf).strip()

    items: list = []
    for line in sections.get("Items", "").split("\n"):
        m = re.match(r"^\d+\.\s+(.+)$", line.strip())
        if m:
            items.append(parse_item(m.group(1)))

    return {
        "title": _strip_bold(sections.get("Title", "").strip()),
        "layout": sections.get("Layout", "Title and Content (common)").strip(),
        "items": items,
        "speaker_notes": sections.get("Speaker notes", ""),
    }


def parse_item(text: str) -> dict:
    """Parse one of:
       - 'Year | Label | Body'  (timeline items, with the | separator)
       - '**Label** — Body'     (block items)
       - 'Label'                (just a label)
    """
    parts = [p.strip() for p in text.split("|")]
    if len(parts) >= 3:
        return {
            "year": _strip_bold(parts[0]),
            "label": _strip_bold(parts[1]),
            "body": parts[2],
        }
    if " — " in text:
        label, body = text.split(" — ", 1)
        return {"year": None, "label": _strip_bold(label.strip()), "body": body.strip()}
    return {"year": None, "label": _strip_bold(text.strip()), "body": ""}


def _strip_bold(text: str) -> str:
    text = text.strip()
    if text.startswith("**") and text.endswith("**"):
        return text[2:-2]
    return text


# --------------------------------------------------------------------------
# Inline markdown -> runs
# --------------------------------------------------------------------------

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")


def add_md_runs(paragraph, text: str):
    """Render markdown-styled text as runs. We do NOT override font size,
    color, bold, or italic at the run level — every formatting attribute is
    inherited from the placeholder's text style defined in the slide layout /
    master, except where the markdown itself dictates (e.g. ** -> bold)."""
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.font.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run.text = part[1:-1]
            run.font.italic = True
        else:
            run.text = part


# --------------------------------------------------------------------------
# Template helpers
# --------------------------------------------------------------------------


def find_layout(prs, name: str):
    for layout in prs.slide_masters[0].slide_layouts:
        if layout.name == name:
            return layout
    raise SystemExit(f"Layout {name!r} not found in master 0")


def get_ph(slide, idx: int):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def write_text(ph, text: str):
    """Replace placeholder text. Inherits all styling from the layout/master."""
    if ph is None or text is None:
        return
    tf = ph.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    add_md_runs(p, text)


def set_title(slide, text: str):
    write_text(get_ph(slide, 0), text)


def set_speaker_notes(slide, notes_text: str):
    if not notes_text:
        return
    notes = slide.notes_slide.notes_text_frame
    notes.clear()
    paragraphs = [p for p in notes_text.split("\n\n") if p.strip()]
    for i, para in enumerate(paragraphs):
        p = notes.paragraphs[0] if i == 0 else notes.add_paragraph()
        run = p.add_run()
        run.text = para.strip()
        run.font.size = Pt(11)


# --------------------------------------------------------------------------
# Per-layout renderers
# --------------------------------------------------------------------------


def render_three_content_blocks(prs, parsed):
    """Layout 'Three Content Blocks': 3 columns, each with subtitle + body."""
    slide = prs.slides.add_slide(find_layout(prs, "Three Content Blocks"))
    set_title(slide, parsed["title"])
    # Subtitles 38/39/40, bodies 10/36/37 (left, middle, right)
    sub_idxs = [38, 39, 40]
    body_idxs = [10, 36, 37]
    for i, item in enumerate(parsed["items"][:3]):
        write_text(get_ph(slide, sub_idxs[i]), item["label"])
        write_text(get_ph(slide, body_idxs[i]), item["body"])
    set_speaker_notes(slide, parsed["speaker_notes"])
    return slide


def render_5_steps(prs, parsed):
    """Layout '5 steps w icons': 5 columns; each has icon (left empty), heading, body."""
    slide = prs.slides.add_slide(find_layout(prs, "5 steps w icons"))
    set_title(slide, parsed["title"])
    # Per layout inspection, ordered left-to-right by x position:
    head_idxs = [22, 31, 34, 37, 69]   # heading row at y≈4.69
    body_idxs = [2, 72, 70, 73, 74]    # body row at y≈5.16
    # Picture placeholders 65, 66, 67, 68, 71 left untouched (user can fill icons)
    for i, item in enumerate(parsed["items"][:5]):
        write_text(get_ph(slide, head_idxs[i]), item["label"])
        write_text(get_ph(slide, body_idxs[i]), item["body"])
    set_speaker_notes(slide, parsed["speaker_notes"])
    return slide


def render_four_content_blocks(prs, parsed):
    """Layout 'Four Content Blocks': 2x2 grid. Reading order top-left → top-right → bottom-left → bottom-right."""
    slide = prs.slides.add_slide(find_layout(prs, "Four Content Blocks"))
    set_title(slide, parsed["title"])
    # Position-mapped:
    # TL: subtitle 50, body 46  | TR: subtitle 52, body 48
    # BL: subtitle 51, body 47  | BR: subtitle 53, body 49
    sub_idxs = [50, 52, 51, 53]
    body_idxs = [46, 48, 47, 49]
    for i, item in enumerate(parsed["items"][:4]):
        write_text(get_ph(slide, sub_idxs[i]), item["label"])
        write_text(get_ph(slide, body_idxs[i]), item["body"])
    set_speaker_notes(slide, parsed["speaker_notes"])
    return slide


def render_six_content_blocks(prs, parsed):
    """Layout 'Six Content Blocks': 2x3 grid. Reading order: TL, TM, TR, BL, BM, BR."""
    slide = prs.slides.add_slide(find_layout(prs, "Six Content Blocks"))
    set_title(slide, parsed["title"])
    # TL: 50/94 | TM: 96/90 | TR: 98/91
    # BL: 51/95 | BM: 97/92 | BR: 99/93
    sub_idxs = [50, 96, 98, 51, 97, 99]
    body_idxs = [94, 90, 91, 95, 92, 93]
    for i, item in enumerate(parsed["items"][:6]):
        write_text(get_ph(slide, sub_idxs[i]), item["label"])
        write_text(get_ph(slide, body_idxs[i]), item["body"])
    set_speaker_notes(slide, parsed["speaker_notes"])
    return slide


def render_timeline_5(prs, parsed):
    """Layout 'Timeline (5 years)': 5 columns, each with year, heading, body."""
    slide = prs.slides.add_slide(find_layout(prs, "Timeline (5 years)"))
    set_title(slide, parsed["title"])
    # Left-to-right by x position:
    year_idxs = [50, 51, 52, 53, 80]   # year markers along the bottom
    head_idxs = [22, 73, 70, 67, 79]   # column headings at top
    body_idxs = [34, 77, 76, 75, 81]   # column body
    for i, item in enumerate(parsed["items"][:5]):
        if item.get("year"):
            write_text(get_ph(slide, year_idxs[i]), item["year"])
        write_text(get_ph(slide, head_idxs[i]), item["label"])
        write_text(get_ph(slide, body_idxs[i]), item["body"])
    set_speaker_notes(slide, parsed["speaker_notes"])
    return slide


LAYOUT_RENDERERS = {
    "Three Content Blocks": render_three_content_blocks,
    "5 steps w icons": render_5_steps,
    "Four Content Blocks": render_four_content_blocks,
    "Six Content Blocks": render_six_content_blocks,
    "Timeline (5 years)": render_timeline_5,
}


# --------------------------------------------------------------------------
# Slice into the right position
# --------------------------------------------------------------------------


def reorder_slides(prs, *, placeholder_idx_0based: int, my_slide_count: int):
    """Move the last `my_slide_count` slides to replace the placeholder."""
    sld_id_lst = prs.slides._sldIdLst
    sld_ids = list(sld_id_lst)
    total = len(sld_ids)
    my_first = total - my_slide_count

    new_order = []
    for i in range(total):
        if i < placeholder_idx_0based:
            new_order.append(sld_ids[i])
        elif i == placeholder_idx_0based:
            new_order.extend(sld_ids[my_first:my_first + my_slide_count])
            # placeholder is dropped
        elif i < my_first:
            new_order.append(sld_ids[i])

    for sld_id in sld_ids:
        sld_id_lst.remove(sld_id)
    for sld_id in new_order:
        sld_id_lst.append(sld_id)


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------


def build_deck():
    if not TEMPLATE.exists():
        raise SystemExit(f"Template not found: {TEMPLATE}")
    if not DECK_DIR.exists():
        raise SystemExit(f"Deck directory not found: {DECK_DIR}")

    prs = Presentation(str(TEMPLATE))
    print(f"Loaded template: {TEMPLATE.name} ({len(prs.slides)} slides)")

    md_files = sorted(p for p in DECK_DIR.glob("*.md") if re.match(r"^\d+-", p.name))
    if not md_files:
        raise SystemExit(f"No slide markdown files found in {DECK_DIR}")

    for md in md_files:
        parsed = parse_compact(md)
        layout_name = parsed["layout"]
        renderer = LAYOUT_RENDERERS.get(layout_name)
        if renderer is None:
            raise SystemExit(f"No renderer registered for layout {layout_name!r} (file: {md.name})")
        renderer(prs, parsed)
        print(f"  rendered: {md.name}  →  {layout_name}")

    reorder_slides(
        prs,
        placeholder_idx_0based=PLACEHOLDER_SLIDE_NUMBER - 1,
        my_slide_count=len(md_files),
    )

    prs.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    build_deck()
